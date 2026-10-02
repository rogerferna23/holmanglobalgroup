#!/usr/bin/env python3
"""
Sistema visual de HGG para documentos.

Los tokens y las piezas de composición salen del sitio (`src/styles/main.css` y
`sistema-visual-hgg.html`): fondo azul-noche, dorado como acento escaso, tarjetas
de superficie con filete tenue, Questrial en titulares y Josefin Sans en el resto.

Lo importan los generadores de documentos de marca (actas de sesión, plan de
lectura). Si cambia la marca, cambia aquí y cambia en todos.
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

# ---------------------------------------------------------------------------
# Sistema de diseño — tokens :root del sitio
# ---------------------------------------------------------------------------

BG = "#0B1016"          # fondo del sitio
BG_2 = "#0E141C"        # tarjetas y superficies
ORO = "#F0B800"         # acento único de marca — escaso, nunca superficie
BLANCO = "#FFFFFF"      # titulares y texto fuerte
TENUE = "#B8BEC7"       # párrafos
TENUE_2 = "#6B7380"     # pies y metadatos

# Filetes: blanco al 8 %. En PDF van con alfa real; en Word se aplanan sobre el
# fondo correspondiente, porque los bordes de Word no admiten transparencia.
FILETE_ALFA = 0.08
FILETE_SOBRE_BG2 = "#21272E"
FILETE_SOBRE_BG = "#1E2329"

GRANO_OPACIDAD = 0.05   # el sitio usa 0.06 en modo overlay

FUENTE_TITULO = "Questrial"
FUENTE_CUERPO = "Josefin Sans"

# La eyebrow del sitio lleva tracking 0.32em. En un documento no hay
# letter-spacing, así que se abre con espacios finos (U+2009). Entre palabras
# hace falta un espacio duro (U+00A0): reportlab colapsa los finos repetidos,
# de modo que cuatro seguidos se ven igual que uno.
FINO = " "
HUECO = " " + FINO

RAIZ = Path(__file__).resolve().parent
LOGO_HGG = RAIZ / "assets" / "logo-hgg.png"
LOGO_ELEFANTE = RAIZ / "assets" / "logo-elefante.png"
GRANO = RAIZ / "assets" / "grano.png"
DIR_FUENTES = Path.home() / "Library" / "Fonts"

MARCA = "Holman Global Group"
TAGLINE = "Corazón de Elefante: Eco, Fuego y Huella"
PIE = "Holman Global Group · Sentido, Marca y Sistema"

MARGEN_X = 58.0
MARGEN_Y = 46.0
ALTO_CIERRE = 66.0      # filete + elefante + firma al pie de la última página


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def _hex(color: str) -> str:
    return color.lstrip("#")


def eyebrow(texto: str) -> str:
    """Mayúsculas con tracking abierto, al estilo de la etiqueta de sección.

    Las letras se separan con espacio fino y las palabras con HUECO; sin eso el
    tracking las funde en una sola mancha.
    """
    return HUECO.join(FINO.join(p) for p in texto.upper().split(" "))


def escapar(texto: str) -> str:
    """Neutraliza los caracteres que reportlab lee como marcado.

    Los `Paragraph` de reportlab parsean XML, así que un "S&P 500" sin escapar
    sale como "S&P; 500". Todo texto que venga de fuera pasa por aquí.
    """
    return (texto.replace("&", "&amp;")
                 .replace("<", "&lt;")
                 .replace(">", "&gt;"))


def slug(texto: str) -> str:
    """Convierte un texto en un fragmento de nombre de archivo seguro."""
    limpio = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    limpio = re.sub(r"[^\w\s-]", "", limpio).strip()
    return re.sub(r"[\s_-]+", "_", limpio) or "sesion"


def parrafos(texto: str) -> list[str]:
    """Divide un bloque de texto en párrafos por líneas en blanco."""
    if not texto:
        return []
    return [p.strip() for p in re.split(r"\n\s*\n", texto.strip()) if p.strip()]


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------

def registrar_fuentes():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    archivos = {
        "Questrial": "Questrial-Regular.ttf",
        "Josefin": "JosefinSans-Regular.ttf",
        "Josefin-Light": "JosefinSans-Light.ttf",
        "Josefin-Medium": "JosefinSans-Medium.ttf",
        "Josefin-SemiBold": "JosefinSans-SemiBold.ttf",
        "Josefin-Bold": "JosefinSans-Bold.ttf",
        "Josefin-Italic": "JosefinSans-Italic.ttf",
    }
    faltantes = [a for a in archivos.values() if not (DIR_FUENTES / a).exists()]
    if faltantes:
        raise SystemExit(
            "Faltan fuentes de marca en ~/Library/Fonts: " + ", ".join(faltantes)
        )
    for alias, archivo in archivos.items():
        pdfmetrics.registerFont(TTFont(alias, str(DIR_FUENTES / archivo)))
    pdfmetrics.registerFontFamily(
        "Josefin", normal="Josefin", bold="Josefin-SemiBold",
        italic="Josefin-Italic", boldItalic="Josefin-Bold",
    )


def _filete(alfa=FILETE_ALFA):
    from reportlab.lib.colors import Color

    return Color(1, 1, 1, alpha=alfa)


def _fondo(canvas, doc):
    """Fondo azul-noche y grano de película, en ese orden."""
    from reportlab.lib.colors import HexColor

    ancho, alto = doc.pagesize
    canvas.saveState()
    canvas.setFillColor(HexColor(BG))
    canvas.rect(0, 0, ancho, alto, stroke=0, fill=1)
    if GRANO.exists():
        canvas.setFillAlpha(GRANO_OPACIDAD)
        lado = 44.0
        y = 0.0
        while y < alto:
            x = 0.0
            while x < ancho:
                canvas.drawImage(str(GRANO), x, y, width=lado, height=lado,
                                 mask="auto")
                x += lado
            y += lado
        canvas.setFillAlpha(1)
    canvas.restoreState()


def _cierre(estado, alto_marco):
    """Bloque de cierre que ocupa lo que quede de página y se pinta al pie.

    Reservar la banda por adelantado desperdiciaba espacio en todas las páginas
    menos la última; empujar con un flowable invisible dejaba el cierre arriba
    cuando ya no cabía. Este se mide contra el hueco disponible: si no le
    alcanza, pide un punto de más y salta de página por sí solo.

    Cuando cae en una página recién abierta anota `estado["solo"]`, para que
    `generar_pdf` recomponga reservando la banda y no quede una hoja con nada
    más que el elefante.
    """
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import Flowable

    class Cierre(Flowable):
        def wrap(self, ancho_disp, alto_disp):
            self.width = ancho_disp
            if alto_disp >= alto_marco - 1:
                estado["solo"] = True
            self.height = alto_disp if alto_disp >= ALTO_CIERRE else alto_disp + 1
            return self.width, self.height

        def draw(self):
            c = self.canv
            lado = 32.0
            centro = self.width / 2
            c.setStrokeColor(_filete())
            c.setLineWidth(0.5)
            c.line(0, 62, self.width, 62)
            c.drawImage(str(LOGO_ELEFANTE), centro - lado / 2, 20,
                        width=lado, height=lado, mask="auto")
            c.setFont("Josefin-Light", 8.5)
            c.setFillColor(HexColor(TENUE_2))
            c.drawCentredString(centro, 7, PIE)

    return Cierre()


def _regla(ancho, espacio_antes=0, espacio_despues=0):
    """Filete horizontal de una hairline de grosor."""
    from reportlab.platypus import Table, TableStyle

    r = Table([[""]], colWidths=[ancho], rowHeights=[0.5])
    r.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 0.5, _filete()),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), espacio_antes),
        ("BOTTOMPADDING", (0, 0), (-1, -1), espacio_despues),
    ]))
    return r



# ---------------------------------------------------------------------------
# DOCX (python-docx)
#
# Word no admite esquinas redondeadas, bordes translúcidos ni grano: las tarjetas
# se aplanan a un relleno sólido con filete opaco equivalente.
# ---------------------------------------------------------------------------

def _xml(tag: str, **attrs):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), v)
    return el


def _fondo_docx(documento):
    """Inyecta <w:background> y <w:displayBackgroundShape>, que python-docx no expone."""
    from docx.oxml.ns import qn

    documento.element.insert(0, _xml("w:background", color=_hex(BG)))
    settings = documento.settings.element
    if settings.find(qn("w:displayBackgroundShape")) is None:
        settings.insert(0, _xml("w:displayBackgroundShape"))


def _sombrear(celda, color: str):
    celda._tc.get_or_add_tcPr().append(
        _xml("w:shd", val="clear", color="auto", fill=_hex(color))
    )


def _bordes(tabla, color: str | None):
    """Filete uniforme, o ninguno si `color` es None."""
    bordes = _xml("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if color is None:
            bordes.append(_xml(f"w:{lado}", val="none", sz="0", space="0",
                               color="auto"))
        else:
            bordes.append(_xml(f"w:{lado}", val="single", sz="4", space="0",
                               color=_hex(color)))
    tabla._tbl.tblPr.append(bordes)


def _anchos_fijos(tabla, proporciones, ancho_total):
    """Fija el ancho de columna en Word, que de lo contrario reparte a su antojo."""
    from docx.shared import Inches

    from docx.oxml.ns import qn

    # `autofit = False` ya escribe <w:tblLayout w:type="fixed">; añadir otro
    # deja dos y Word descarta el bloque entero de propiedades.
    tabla.autofit = False
    # Sin un tblW explícito Word encoge la tabla al ancho de su contenido y los
    # anchos de celda quedan en nada.
    tblw = tabla._tbl.tblPr.find(qn("w:tblW"))
    if tblw is not None:
        tblw.set(qn("w:w"), str(int(ancho_total * 1440)))
        tblw.set(qn("w:type"), "dxa")
    anchos = [Inches(ancho_total * p) for p in proporciones]
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos):
            celda.width = ancho
    for columna, ancho in zip(tabla.columns, anchos):
        columna.width = ancho


def _run(parrafo, texto, *, fuente=FUENTE_CUERPO, tam=10.5, color=TENUE,
         negrita=False):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    r = parrafo.add_run(texto)
    r.font.size = Pt(tam)
    r.font.color.rgb = RGBColor.from_string(_hex(color))
    r.bold = negrita
    r.font.name = fuente
    rfonts = r._element.get_or_add_rPr().find(qn("w:rFonts"))
    if rfonts is not None:
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), fuente)
    return r


def _regla_docx(doc, ancho_util):
    """Filete horizontal: una tabla de una celda con borde solo abajo."""
    from docx.shared import Pt

    tabla = doc.add_table(rows=1, cols=1)
    _bordes(tabla, None)
    _anchos_fijos(tabla, (1.0,), ancho_util)
    celda = tabla.rows[0].cells[0]
    borde = _xml("w:tcBorders")
    borde.append(_xml("w:bottom", val="single", sz="4", space="0",
                      color=_hex(FILETE_SOBRE_BG)))
    celda._tc.get_or_add_tcPr().append(borde)
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    _run(p, "", tam=1)
    return tabla



# ---------------------------------------------------------------------------
# Cabecera de marca — la misma en todos los documentos
# ---------------------------------------------------------------------------

def cabecera_pdf(ancho_util):
    """Logo, nombre, tagline y filete. Devuelve la lista de flowables."""
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, Spacer, Table, TableStyle

    est_marca = ParagraphStyle(
        "marca", fontName="Questrial", fontSize=13, leading=16,
        textColor=HexColor(BLANCO),
    )
    est_tagline = ParagraphStyle(
        "tagline", fontName="Josefin-Light", fontSize=9.5, leading=13,
        textColor=HexColor(TENUE_2),
    )
    logo = Image(str(LOGO_HGG), width=0.68 * inch, height=0.68 * inch)
    tabla = Table(
        [[logo, [Paragraph(MARCA, est_marca), Paragraph(TAGLINE, est_tagline)]]],
        colWidths=[0.82 * inch, ancho_util - 0.82 * inch],
    )
    tabla.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    return [tabla, Spacer(1, 10), _regla(ancho_util)]


def cabecera_docx(doc, ancho_util):
    """Equivalente en Word de `cabecera_pdf`."""
    from docx.shared import Inches, Pt

    cab = doc.add_table(rows=1, cols=2)
    _bordes(cab, None)
    _anchos_fijos(cab, (0.12, 0.88), ancho_util)
    celda_logo, celda_texto = cab.rows[0].cells
    celda_logo.paragraphs[0].add_run().add_picture(str(LOGO_HGG), width=Inches(0.68))
    p_marca = celda_texto.paragraphs[0]
    p_marca.paragraph_format.space_after = Pt(0)
    _run(p_marca, MARCA, fuente=FUENTE_TITULO, tam=13, color=BLANCO)
    p_tag = celda_texto.add_paragraph()
    p_tag.paragraph_format.space_after = Pt(0)
    _run(p_tag, TAGLINE, tam=9.5, color=TENUE_2)
    _regla_docx(doc, ancho_util)
    return cab
