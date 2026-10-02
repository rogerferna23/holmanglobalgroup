#!/usr/bin/env python3
"""
Generador de actas de Sesión 1 a 1 HGG.

Toma un JSON con el contenido de la sesión y produce el acta en PDF y DOCX con el
sistema visual de la marca, que vive en `marca_hgg.py`.

Uso:
    python3 generar_sesion.py sesion.json [--salida DIRECTORIO] [--solo pdf|docx]

Esquema del JSON (todos los campos son opcionales salvo `coachee`):

{
  "coachee": "Valeria Aponte",
  "fecha": "28.11.25",
  "tipo_sesion": "Coaching",
  "contacto": "+57 301 4431676",
  "objetivo": "Sesión de autoconocimiento",
  "resumen": "Párrafo o varios separados por \\n\\n",
  "hallazgos": ["...", "..."],
  "plan_accion": [
    {"tarea": "...", "responsable": "...", "fecha_limite": "...", "metrica": "..."}
  ]
}
"""

from __future__ import annotations

import argparse
import io
import json
import sys
from pathlib import Path

from marca_hgg import (
    ALTO_CIERRE, PIE, RAIZ, BG_2, BLANCO, FUENTE_CUERPO, FUENTE_TITULO, LOGO_ELEFANTE,
    LOGO_HGG, MARCA, MARGEN_X, MARGEN_Y, ORO, TAGLINE, TENUE, TENUE_2,
    FILETE_SOBRE_BG2, _anchos_fijos, _bordes, _cierre, _filete, _fondo,
    _fondo_docx, _regla, _run, _sombrear, cabecera_docx, cabecera_pdf,
    escapar, eyebrow, parrafos, registrar_fuentes, slug,
)

TITULO_DOC = "Sesión 1 a 1 HGG"


def cargar_sesion(ruta: Path) -> dict:
    datos = json.loads(ruta.read_text(encoding="utf-8"))
    if not datos.get("coachee"):
        raise SystemExit("El JSON debe incluir al menos el campo 'coachee'.")
    for campo in ("fecha", "contacto", "objetivo", "resumen"):
        datos.setdefault(campo, "")
    if not datos.get("tipo_sesion"):
        datos["tipo_sesion"] = "Coaching"
    datos.setdefault("hallazgos", [])
    datos.setdefault("plan_accion", [])
    return datos



# -------------------------------------------------------------------------
# PDF
# -------------------------------------------------------------------------

def generar_pdf(sesion: dict, destino: Path) -> Path:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        BaseDocTemplate, CondPageBreak, Frame, Image, PageTemplate,
        Paragraph, Spacer, Table, TableStyle,
    )

    registrar_fuentes()
    ancho_util = letter[0] - 2 * MARGEN_X

    est_eyebrow = ParagraphStyle(
        "eyebrow", fontName="Josefin-SemiBold", fontSize=7.5, leading=11,
        textColor=HexColor(ORO), spaceBefore=20, spaceAfter=9, keepWithNext=1,
    )
    est_titulo = ParagraphStyle(
        "titulo", fontName="Questrial", fontSize=26, leading=29,
        textColor=HexColor(BLANCO), spaceBefore=4, spaceAfter=18,
    )
    est_cuerpo = ParagraphStyle(
        "cuerpo", fontName="Josefin-Light", fontSize=10.5, leading=17.5,
        textColor=HexColor(TENUE), spaceAfter=11,
        allowWidows=0, allowOrphans=0,
    )
    est_vineta = ParagraphStyle(
        "vineta", parent=est_cuerpo, leftIndent=15, bulletIndent=2,
        bulletFontName="Josefin", bulletFontSize=10.5, spaceAfter=13,
        bulletColor=HexColor(ORO),
    )
    est_dato_etiqueta = ParagraphStyle(
        "datoEtiqueta", fontName="Josefin-SemiBold", fontSize=6.5, leading=10,
        textColor=HexColor(TENUE_2), spaceAfter=2,
    )
    est_dato_valor = ParagraphStyle(
        "datoValor", fontName="Josefin", fontSize=11, leading=14,
        textColor=HexColor(BLANCO),
    )
    # El plan es lo único que puede partirse entre páginas: su etiqueta no se
    # pega a la tabla, o un plan largo dejaría media página en blanco.
    est_eyebrow_plan = ParagraphStyle("eyebrowPlan", parent=est_eyebrow,
                                      keepWithNext=0)
    est_col = ParagraphStyle(
        "col", fontName="Josefin-SemiBold", fontSize=6.5, leading=10,
        textColor=HexColor(ORO),
    )
    est_tarea = ParagraphStyle(
        "tarea", fontName="Josefin", fontSize=9.5, leading=14.5,
        textColor=HexColor(BLANCO),
    )
    est_celda = ParagraphStyle(
        "celda", fontName="Josefin-Light", fontSize=9.5, leading=14.5,
        textColor=HexColor(TENUE),
    )

    tarjeta_base = [
        ("BACKGROUND", (0, 0), (-1, -1), HexColor(BG_2)),
        ("BOX", (0, 0), (-1, -1), 0.5, _filete()),
        ("ROUNDEDCORNERS", [10, 10, 10, 10]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 13),
        ("RIGHTPADDING", (0, 0), (-1, -1), 13),
        ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
    ]

    def construir_cuerpo():
        """Devuelve un flujo nuevo en cada llamada.

        reportlab consume los flowables al partirlos entre páginas, así que la
        segunda pasada no puede reutilizar los de la primera.
        """
        flujo = []

        flujo += cabecera_pdf(ancho_util)

        # --- Titular: eyebrow dorada + display en blanco ------------------------
        etiqueta = "Acta de sesión"
        if sesion["fecha"]:
            etiqueta += f" · {sesion['fecha']}"
        flujo += [
            Paragraph(eyebrow(etiqueta), est_eyebrow),
            Paragraph(TITULO_DOC, est_titulo),
        ]

        # --- Tarjeta de datos ---------------------------------------------------
        def dato(nombre, valor):
            return [Paragraph(eyebrow(nombre), est_dato_etiqueta),
                    Paragraph(escapar(valor) or "—", est_dato_valor)]

        tarjeta = Table(
            [[dato("Coachee", sesion["coachee"]), dato("Fecha", sesion["fecha"])],
             [dato("Tipo de sesión", sesion["tipo_sesion"]),
              dato("Contacto", sesion["contacto"])]],
            colWidths=[ancho_util / 2] * 2,
        )
        tarjeta.setStyle(TableStyle(tarjeta_base + [
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, _filete()),
            ("LINEAFTER", (0, 0), (0, -1), 0.5, _filete()),
        ]))
        flujo.append(tarjeta)

        # --- Objetivo y resumen -------------------------------------------------
        for titulo, campo in (("Objetivo de la sesión", "objetivo"),
                              ("Resumen de la sesión", "resumen")):
            if sesion[campo]:
                flujo.append(Paragraph(eyebrow(titulo), est_eyebrow))
                for p in parrafos(sesion[campo]):
                    flujo.append(Paragraph(escapar(p), est_cuerpo))

        # --- Hallazgos ----------------------------------------------------------
        if sesion["hallazgos"]:
            flujo.append(Paragraph(eyebrow("Hallazgos"), est_eyebrow))
            for h in sesion["hallazgos"]:
                flujo.append(Paragraph(escapar(h), est_vineta, bulletText="•"))

        # --- Plan de acción -----------------------------------------------------
        if sesion["plan_accion"]:
            flujo.append(CondPageBreak(78))
            flujo.append(Paragraph(eyebrow("Plan de acción (SMART)"), est_eyebrow_plan))
            columnas = ["Tarea", "Responsable", "Fecha límite", "Métrica de éxito"]
            filas = [[Paragraph(c.upper(), est_col) for c in columnas]]
            for t in sesion["plan_accion"]:
                filas.append([
                    Paragraph(escapar(t.get("tarea", "")), est_tarea),
                    Paragraph(escapar(t.get("responsable", "")), est_celda),
                    Paragraph(escapar(t.get("fecha_limite", "")), est_celda),
                    Paragraph(escapar(t.get("metrica", "")), est_celda),
                ])
            anchos = [ancho_util * f for f in (0.38, 0.16, 0.20, 0.26)]
            plan = Table(filas, colWidths=anchos, repeatRows=1)
            plan.setStyle(TableStyle(tarjeta_base + [
                ("TOPPADDING", (0, 0), (-1, -1), 10),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
                ("LINEBELOW", (0, 0), (-1, -2), 0.5, _filete()),
            ]))
            flujo.append(plan)
        return flujo


    def documento(salida, reserva):
        doc = BaseDocTemplate(
            salida, pagesize=letter,
            leftMargin=MARGEN_X, rightMargin=MARGEN_X,
            topMargin=MARGEN_Y, bottomMargin=MARGEN_Y + reserva,
            title=f"{TITULO_DOC} — {sesion['coachee']}",
            author=MARCA, subject=TAGLINE,
        )
        marco = Frame(
            doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="cuerpo",
            leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0,
        )
        doc.addPageTemplates([PageTemplate(id="hgg", frames=[marco],
                                           onPage=_fondo)])
        return doc

    # Primera pasada al aire: si el cierre acaba solo en una hoja, se rehace
    # reservando su banda para que el contenido se redistribuya.
    ensayo = documento(io.BytesIO(), 0)
    estado = {}
    ensayo.build(construir_cuerpo() + [_cierre(estado, ensayo.height)])

    reserva = ALTO_CIERRE if estado.get("solo") else 0
    destino.parent.mkdir(parents=True, exist_ok=True)
    final = documento(str(destino), reserva)
    final.build(construir_cuerpo() + [_cierre({}, final.height)])
    return destino


# ---------------------------------------------------------------------------
# DOCX (python-docx)
#
# Word no admite esquinas redondeadas, bordes translúcidos ni grano: las tarjetas
# se aplanan a un relleno sólido con filete opaco equivalente.
# ---------------------------------------------------------------------------

def generar_docx(sesion: dict, destino: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    _fondo_docx(doc)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Pt(MARGEN_X)
    sec.top_margin = sec.bottom_margin = Pt(MARGEN_Y)

    normal = doc.styles["Normal"]
    normal.font.name = FUENTE_CUERPO
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(11)
    normal.paragraph_format.line_spacing = 1.4

    ancho_util = 8.5 - 2 * (MARGEN_X / 72)

    def etiqueta_seccion(titulo: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        _run(p, eyebrow(titulo), tam=7.5, color=ORO, negrita=True)
        return p

    cabecera_docx(doc, ancho_util)

    # --- Titular ------------------------------------------------------------
    etiqueta = "Acta de sesión"
    if sesion["fecha"]:
        etiqueta += f" · {sesion['fecha']}"
    p_eyebrow = doc.add_paragraph()
    p_eyebrow.paragraph_format.space_before = Pt(14)
    p_eyebrow.paragraph_format.space_after = Pt(3)
    _run(p_eyebrow, eyebrow(etiqueta), tam=7.5, color=ORO, negrita=True)

    p_titulo = doc.add_paragraph()
    p_titulo.paragraph_format.space_after = Pt(13)
    _run(p_titulo, TITULO_DOC, fuente=FUENTE_TITULO, tam=26, color=BLANCO)

    # --- Tarjeta de datos ---------------------------------------------------
    tarjeta = doc.add_table(rows=2, cols=2)
    _bordes(tarjeta, FILETE_SOBRE_BG2)
    _anchos_fijos(tarjeta, (0.5, 0.5), ancho_util)
    datos = [[("Coachee", sesion["coachee"]), ("Fecha", sesion["fecha"])],
             [("Tipo de sesión", sesion["tipo_sesion"]),
              ("Contacto", sesion["contacto"])]]
    for fila, contenido in zip(tarjeta.rows, datos):
        for celda, (nombre, valor) in zip(fila.cells, contenido):
            _sombrear(celda, BG_2)
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            _run(p, eyebrow(nombre), tam=6.5, color=TENUE_2, negrita=True)
            pv = celda.add_paragraph()
            pv.paragraph_format.space_after = Pt(0)
            _run(pv, valor or "—", tam=11, color=BLANCO)

    # --- Objetivo y resumen -------------------------------------------------
    for titulo, campo in (("Objetivo de la sesión", "objetivo"),
                          ("Resumen de la sesión", "resumen")):
        if sesion[campo]:
            etiqueta_seccion(titulo)
            for texto in parrafos(sesion[campo]):
                _run(doc.add_paragraph(), texto)

    # --- Hallazgos ----------------------------------------------------------
    if sesion["hallazgos"]:
        etiqueta_seccion("Hallazgos")
        for h in sesion["hallazgos"]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-12)
            p.paragraph_format.space_after = Pt(13)
            _run(p, "•  ", color=ORO)
            _run(p, h)

    # --- Plan de acción -----------------------------------------------------
    if sesion["plan_accion"]:
        etiqueta_seccion("Plan de acción (SMART)")
        columnas = ["Tarea", "Responsable", "Fecha límite", "Métrica de éxito"]
        proporciones = (0.38, 0.16, 0.20, 0.26)
        plan = doc.add_table(rows=1 + len(sesion["plan_accion"]), cols=4)
        _bordes(plan, FILETE_SOBRE_BG2)
        _anchos_fijos(plan, proporciones, ancho_util)
        for i, texto in enumerate(columnas):
            celda = plan.rows[0].cells[i]
            _sombrear(celda, BG_2)
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            _run(p, texto.upper(), tam=6.5, color=ORO, negrita=True)
        for fila, tarea in zip(plan.rows[1:], sesion["plan_accion"]):
            valores = [tarea.get("tarea", ""), tarea.get("responsable", ""),
                       tarea.get("fecha_limite", ""), tarea.get("metrica", "")]
            for i, valor in enumerate(valores):
                celda = fila.cells[i]
                _sombrear(celda, BG_2)
                p = celda.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                _run(p, valor, tam=9.5, color=BLANCO if i == 0 else TENUE)

    # --- Cierre -------------------------------------------------------------
    p_elefante = doc.add_paragraph()
    p_elefante.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_elefante.paragraph_format.space_before = Pt(22)
    p_elefante.paragraph_format.space_after = Pt(4)
    p_elefante.add_run().add_picture(str(LOGO_ELEFANTE), width=Inches(0.44))
    p_pie = doc.add_paragraph()
    p_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p_pie, PIE, tam=8.5, color=TENUE_2)

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return destino

# ---------------------------------------------------------------------------

def main() -> int:
    ap = argparse.ArgumentParser(description="Genera el acta de una sesión 1 a 1 HGG.")
    ap.add_argument("json", type=Path, help="Ruta al JSON de la sesión")
    ap.add_argument("--salida", type=Path, default=RAIZ / "generadas",
                    help="Directorio de salida (por defecto: ./generadas)")
    ap.add_argument("--solo", choices=("pdf", "docx"), default=None,
                    help="Genera solo uno de los dos formatos")
    args = ap.parse_args()

    sesion = cargar_sesion(args.json)
    fecha = slug(sesion["fecha"]) or "sin_fecha"
    base = f"{slug(sesion['coachee'])}_{fecha}_Sesion_1a1_HGG"

    generados = []
    if args.solo != "docx":
        generados.append(generar_pdf(sesion, args.salida / f"{base}.pdf"))
    if args.solo != "pdf":
        generados.append(generar_docx(sesion, args.salida / f"{base}.docx"))

    for ruta in generados:
        print(ruta)
    return 0


if __name__ == "__main__":
    sys.exit(main())
