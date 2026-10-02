#!/usr/bin/env python3
"""
Generador del Plan de Lectura Bíblica de HGG.

Produce en PDF y DOCX el plan de 93 semanas con el sistema visual de la marca
(`marca_hgg.py`), precedido por el método de lectura: leer, subrayar, investigar,
meditar, anotar y aplicar.

Uso:
    python3 generar_plan_lectura.py [--para NOMBRE] [--salida DIR] [--solo pdf|docx]

Los datos del plan viven en `datos/plan-lectura-biblica.json`.
"""

from __future__ import annotations

import argparse
import io
import json
import sys
from pathlib import Path

from marca_hgg import (
    ALTO_CIERRE, BG_2, BLANCO, FILETE_SOBRE_BG2, FUENTE_CUERPO, FUENTE_TITULO,
    MARCA, MARGEN_X, MARGEN_Y, ORO, PIE, RAIZ, TAGLINE, TENUE, TENUE_2,
    _anchos_fijos, _bordes, _cierre, _filete, _fondo, _fondo_docx, _run,
    _sombrear, cabecera_docx, cabecera_pdf, escapar, eyebrow, registrar_fuentes, slug,
)

DATOS = RAIZ / "datos" / "plan-lectura-biblica.json"

TITULO_DOC = "Plan de Lectura Bíblica"

ENTRADA = (
    "Noventa y tres semanas para leer la Biblia entera. Cada día de la semana "
    "tiene su propia lectura, y cada una avanza por una parte distinta: relatos, "
    "historia, Salmos, poesía y proverbios, profetas, Evangelios y Hechos, y "
    "cartas. Siete caminos que llegan juntos al final."
)

PASOS = [
    ("Lee la porción del día",
     "Busca tu semana en la tabla y lee lo que corresponde a hoy. A la misma "
     "hora, todos los días: lo que sostiene esto es el hábito, y el hábito se "
     "construye con la cita cumplida, no con las ganas."),
    ("Subraya mientras lees",
     "En JW Library, subraya los textos que más te llamen la atención. No los "
     "elijas por importantes: elígelos porque te movieron algo."),
    ("Escoge dos o tres versículos por capítulo",
     "Al terminar cada capítulo, quédate con dos o tres versículos e investiga "
     "sobre ellos: quién lo dijo, a quién, qué estaba pasando alrededor."),
    ("Medita con las tres preguntas",
     "Sobre esos versículos, hazte las tres preguntas de abajo. Ahí es donde la "
     "lectura deja de ser información y se vuelve tuya."),
    ("Escribe tu comentario",
     "Crea una nota en ese versículo dentro de JW Library y escribe lo que "
     "aprendiste, con tus palabras. Lo que no se escribe se olvida."),
    ("Llévalo al día",
     "Durante la jornada, vuelve sobre esos puntos y piensa cómo aplicarlos en "
     "lo que estás viviendo hoy."),
]

PREGUNTAS = [
    "¿Qué me enseña esto sobre Jehová?",
    "¿Cómo puedo aplicar esto para ayudar a los demás?",
    "¿Qué me está queriendo decir Jehová en este texto?",
]

# La semana ocupa poco; los siete días se reparten el resto por igual.
PROPORCIONES = (0.078,) + (0.1317,) * 7


def cargar_plan() -> dict:
    if not DATOS.exists():
        raise SystemExit(f"No encuentro los datos del plan en {DATOS}")
    return json.loads(DATOS.read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def generar_pdf(plan: dict, para: str, destino: Path) -> Path:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (
        BaseDocTemplate, CondPageBreak, Frame, PageTemplate, Paragraph,
        Table, TableStyle,
    )

    registrar_fuentes()
    ancho_util = letter[0] - 2 * MARGEN_X

    est_eyebrow = ParagraphStyle(
        "eyebrow", fontName="Josefin-SemiBold", fontSize=7.5, leading=11,
        textColor=HexColor(ORO), spaceBefore=20, spaceAfter=9, keepWithNext=1,
    )
    est_titulo = ParagraphStyle(
        "titulo", fontName="Questrial", fontSize=26, leading=29,
        textColor=HexColor(BLANCO), spaceBefore=4, spaceAfter=14,
    )
    est_entrada = ParagraphStyle(
        "entrada", fontName="Josefin-Light", fontSize=11.5, leading=19,
        textColor=HexColor(TENUE), spaceAfter=4,
        allowWidows=0, allowOrphans=0,
    )
    est_paso_titulo = ParagraphStyle(
        "pasoTitulo", fontName="Josefin", fontSize=10.5, leading=15,
        textColor=HexColor(BLANCO), spaceAfter=2,
    )
    est_paso_texto = ParagraphStyle(
        "pasoTexto", fontName="Josefin-Light", fontSize=10, leading=15.5,
        textColor=HexColor(TENUE), allowWidows=0, allowOrphans=0,
    )
    est_numero = ParagraphStyle(
        "numero", fontName="Questrial", fontSize=17, leading=19,
        textColor=HexColor(ORO),
    )
    est_pregunta = ParagraphStyle(
        "pregunta", fontName="Josefin", fontSize=11, leading=17,
        textColor=HexColor(BLANCO), leftIndent=14, bulletIndent=2,
        bulletFontName="Josefin", bulletFontSize=11,
        bulletColor=HexColor(ORO), spaceAfter=6,
    )
    # La tabla del plan es enorme y tiene que poder partirse: si su etiqueta se
    # le pega, arrastra las 93 semanas a la página siguiente y deja un hueco.
    est_eyebrow_tabla = ParagraphStyle("eyebrowTabla", parent=est_eyebrow,
                                       keepWithNext=0)
    est_col = ParagraphStyle(
        "col", fontName="Josefin-SemiBold", fontSize=6.5, leading=10,
        textColor=HexColor(ORO),
    )
    est_semana = ParagraphStyle(
        "semana", fontName="Josefin", fontSize=7, leading=10,
        textColor=HexColor(TENUE_2),
    )
    est_lectura = ParagraphStyle(
        "lectura", fontName="Josefin-Light", fontSize=7, leading=10,
        textColor=HexColor(TENUE),
    )

    tarjeta_base = [
        ("BACKGROUND", (0, 0), (-1, -1), HexColor(BG_2)),
        ("BOX", (0, 0), (-1, -1), 0.5, _filete()),
        ("ROUNDEDCORNERS", [10, 10, 10, 10]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 13),
        ("RIGHTPADDING", (0, 0), (-1, -1), 13),
        ("TOPPADDING", (0, 0), (-1, -1), 12),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
    ]

    def construir_cuerpo():
        flujo = list(cabecera_pdf(ancho_util))

        etiqueta = "Plan de lectura · 93 semanas"
        if para:
            etiqueta = f"Para {para} · 93 semanas"
        flujo += [
            Paragraph(eyebrow(etiqueta), est_eyebrow),
            Paragraph(TITULO_DOC, est_titulo),
            Paragraph(ENTRADA, est_entrada),
        ]

        # --- El método -----------------------------------------------------
        flujo.append(Paragraph(eyebrow("Cómo usarlo"), est_eyebrow))
        for i, (titulo, texto) in enumerate(PASOS, start=1):
            paso = Table(
                [[Paragraph(str(i), est_numero),
                  [Paragraph(titulo, est_paso_titulo),
                   Paragraph(texto, est_paso_texto)]]],
                colWidths=[26, ancho_util - 26],
            )
            paso.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, 0), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (0, 0), 8),
                ("RIGHTPADDING", (1, 0), (1, 0), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 13),
            ]))
            flujo.append(paso)

        # --- Las tres preguntas --------------------------------------------
        flujo.append(Paragraph(eyebrow("Las tres preguntas"), est_eyebrow))
        tarjeta = Table(
            [[[Paragraph(q, est_pregunta, bulletText="•") for q in PREGUNTAS]]],
            colWidths=[ancho_util],
        )
        tarjeta.setStyle(TableStyle(tarjeta_base))
        flujo.append(tarjeta)

        # --- El plan --------------------------------------------------------
        flujo.append(CondPageBreak(90))
        flujo.append(Paragraph(eyebrow("El plan, semana a semana"),
                               est_eyebrow_tabla))
        filas = [[Paragraph(c.upper(), est_col) for c in plan["columnas"]]]
        for semana in plan["semanas"]:
            filas.append(
                [Paragraph(semana[0], est_semana)]
                + [Paragraph(c, est_lectura) for c in semana[1:]]
            )
        anchos = [ancho_util * f for f in PROPORCIONES]
        tabla = Table(filas, colWidths=anchos, repeatRows=1)
        tabla.setStyle(TableStyle(tarjeta_base + [
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LINEBELOW", (0, 0), (-1, -2), 0.5, _filete()),
        ]))
        flujo.append(tabla)
        return flujo

    def documento(salida, reserva):
        doc = BaseDocTemplate(
            salida, pagesize=letter,
            leftMargin=MARGEN_X, rightMargin=MARGEN_X,
            topMargin=MARGEN_Y, bottomMargin=MARGEN_Y + reserva,
            title=TITULO_DOC, author=MARCA, subject=TAGLINE,
        )
        marco = Frame(
            doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="cuerpo",
            leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0,
        )
        doc.addPageTemplates([PageTemplate(id="hgg", frames=[marco],
                                           onPage=_fondo)])
        return doc

    ensayo = documento(io.BytesIO(), 0)
    estado = {}
    ensayo.build(construir_cuerpo() + [_cierre(estado, ensayo.height)])

    reserva = ALTO_CIERRE if estado.get("solo") else 0
    destino.parent.mkdir(parents=True, exist_ok=True)
    final = documento(str(destino), reserva)
    final.build(construir_cuerpo() + [_cierre({}, final.height)])
    return destino


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def generar_docx(plan: dict, para: str, destino: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    _fondo_docx(doc)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Pt(MARGEN_X)
    sec.top_margin = sec.bottom_margin = Pt(MARGEN_Y)

    normal = doc.styles["Normal"]
    normal.font.name = FUENTE_CUERPO
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.4

    ancho_util = 8.5 - 2 * (MARGEN_X / 72)

    def etiqueta_seccion(titulo: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        _run(p, eyebrow(titulo), tam=7.5, color=ORO, negrita=True)
        return p

    cabecera_docx(doc, ancho_util)

    etiqueta = "Plan de lectura · 93 semanas"
    if para:
        etiqueta = f"Para {para} · 93 semanas"
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(3)
    _run(p, eyebrow(etiqueta), tam=7.5, color=ORO, negrita=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(13)
    _run(p, TITULO_DOC, fuente=FUENTE_TITULO, tam=26, color=BLANCO)

    _run(doc.add_paragraph(), ENTRADA, tam=11.5, color=TENUE)

    # --- El método ----------------------------------------------------------
    etiqueta_seccion("Cómo usarlo")
    for i, (titulo, texto) in enumerate(PASOS, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        _run(p, f"{i}.  ", fuente=FUENTE_TITULO, tam=12, color=ORO)
        _run(p, titulo, tam=10.5, color=BLANCO)
        pt = doc.add_paragraph()
        pt.paragraph_format.left_indent = Pt(20)
        pt.paragraph_format.space_after = Pt(11)
        _run(pt, texto, tam=10, color=TENUE)

    # --- Las tres preguntas -------------------------------------------------
    etiqueta_seccion("Las tres preguntas")
    tarjeta = doc.add_table(rows=1, cols=1)
    _bordes(tarjeta, FILETE_SOBRE_BG2)
    _anchos_fijos(tarjeta, (1.0,), ancho_util)
    celda = tarjeta.rows[0].cells[0]
    _sombrear(celda, BG_2)
    for i, q in enumerate(PREGUNTAS):
        p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-12)
        p.paragraph_format.space_after = Pt(0 if i == len(PREGUNTAS) - 1 else 6)
        _run(p, "•  ", tam=11, color=ORO)
        _run(p, q, tam=11, color=BLANCO)

    # --- El plan ------------------------------------------------------------
    etiqueta_seccion("El plan, semana a semana")
    tabla = doc.add_table(rows=1 + len(plan["semanas"]), cols=8)
    _bordes(tabla, FILETE_SOBRE_BG2)
    _anchos_fijos(tabla, PROPORCIONES, ancho_util)
    for i, texto in enumerate(plan["columnas"]):
        celda = tabla.rows[0].cells[i]
        _sombrear(celda, BG_2)
        p = celda.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        _run(p, texto.upper(), tam=6.5, color=ORO, negrita=True)
    for fila, semana in zip(tabla.rows[1:], plan["semanas"]):
        for i, valor in enumerate(semana):
            celda = fila.cells[i]
            _sombrear(celda, BG_2)
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            _run(p, valor, tam=7, color=TENUE_2 if i == 0 else TENUE)

    # --- Cierre -------------------------------------------------------------
    from marca_hgg import LOGO_ELEFANTE

    p_elefante = doc.add_paragraph()
    p_elefante.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_elefante.paragraph_format.space_before = Pt(22)
    p_elefante.paragraph_format.space_after = Pt(4)
    p_elefante.add_run().add_picture(str(LOGO_ELEFANTE), width=Inches(0.44))
    p_pie = doc.add_paragraph()
    p_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p_pie, PIE, tam=8.5, color=TENUE_2)

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return destino


# ---------------------------------------------------------------------------

def main() -> int:
    ap = argparse.ArgumentParser(description="Genera el Plan de Lectura Bíblica HGG.")
    ap.add_argument("--para", default="", help="Nombre del coachee, para la etiqueta")
    ap.add_argument("--salida", type=Path, default=RAIZ / "generadas",
                    help="Directorio de salida (por defecto: ./generadas)")
    ap.add_argument("--solo", choices=("pdf", "docx"), default=None,
                    help="Genera solo uno de los dos formatos")
    args = ap.parse_args()

    plan = cargar_plan()
    base = "Plan_Lectura_Biblica_HGG"
    if args.para:
        base = f"{slug(args.para)}_{base}"

    generados = []
    if args.solo != "docx":
        generados.append(generar_pdf(plan, args.para, args.salida / f"{base}.pdf"))
    if args.solo != "pdf":
        generados.append(generar_docx(plan, args.para, args.salida / f"{base}.docx"))

    for ruta in generados:
        print(ruta)
    return 0


if __name__ == "__main__":
    sys.exit(main())
